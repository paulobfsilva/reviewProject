#!/usr/bin/env python3
# scripts/review/build_docx.py
from pathlib import Path
import argparse
import re

from docx import Document
from docx.shared import Pt            # <-- proper import
from docx.enum.text import WD_BREAK

FINAL_DIR  = Path("output/final")
REVIEW_DIR = Path("output/review")
REVIEW_DIR.mkdir(parents=True, exist_ok=True)

def add_para(doc: Document, text: str) -> None:
    """Add a paragraph with 0-pt spacing and 12-pt font."""
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    for run in p.runs:
        run.font.size = Pt(12)

# ---------------------------------------------------------------------
def parse_final_txt(path: Path):
    """
    Return title, place, date, [(num, text)] from a *_final.txt file.
    Joins multi-line numbered blocks into one paragraph.
    """
    lines = [ln.rstrip() for ln in path.read_text(encoding="utf-8").splitlines()
             if ln.strip()]
    if len(lines) < 4:
        raise ValueError(f"{path} needs at least title / place / date / 1 paragraph")

    title, place, date = lines[:3]

    paras, cur_num, buf = [], None, []
    for ln in lines[3:]:
        m = re.match(r"^(\d+)\s+(.*)$", ln)
        if m:
            if cur_num is not None:
                paras.append((cur_num, "\n".join(buf).rstrip()))
            cur_num, buf = int(m.group(1)), [m.group(2)]
        else:
            buf.append(ln)
    if cur_num is not None:
        paras.append((cur_num, "\n".join(buf).rstrip()))
    return title, place, date, paras

# ---------------------------------------------------------------------
def txt_to_docx(sermon_id: str):
    src = FINAL_DIR / f"{sermon_id}_final.txt"
    if not src.exists():
        print(f"[skip] {sermon_id}: no *_final.txt found")
        return

    title, place, date, paras = parse_final_txt(src)
    doc = Document()

    # header lines
    for line in (title, place, date):
        add_para(doc, line)

    # numbered paragraphs with internal line-breaks
    for num, txt in paras:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)

        run_num = p.add_run(f"{num} ")
        run_num.font.size = Pt(12)

        parts = txt.split("\n")
        for i, part in enumerate(parts):
            run = p.add_run(part)
            run.font.size = Pt(12)
            if i != len(parts) - 1:
                run.add_break(WD_BREAK.LINE)

    out_path = REVIEW_DIR / f"{sermon_id}.docx"
    doc.save(out_path)
    print("Saved", out_path)

# ---------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(
        description="Convert *_final.txt files to reviewer-friendly DOCX.")
    ap.add_argument("--sermon-id", help="Convert just this sermon (e.g. 55-0116A)")
    ap.add_argument("--all", action="store_true", help="Convert all *_final.txt files")
    args = ap.parse_args()

    if args.all:
        txt_files = sorted(FINAL_DIR.glob("*_final.txt"))
        if not txt_files:
            print("No *_final.txt files found in", FINAL_DIR)
        for f in txt_files:
            txt_to_docx(f.stem.replace("_final", ""))
        print("Batch done:", len(txt_files), "files processed.")
    elif args.sermon_id:
        txt_to_docx(args.sermon_id)
    else:
        ap.error("Provide --sermon-id or --all")

# ---------------------------------------------------------------------
if __name__ == "__main__":
    main()
